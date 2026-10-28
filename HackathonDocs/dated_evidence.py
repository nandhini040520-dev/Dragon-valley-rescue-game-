from pathlib import Path
import hashlib
import json
import shutil

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent
SOURCE = Path('C:/Users/Karthika S/OneDrive/Pictures/Screenshots')
DESTINATION = ROOT / 'Dated-Asset-Evidence'
HEADING = 'Appendix A. Dated Unity asset evidence'
SHOTS = [
    ('204214', 'Aster dragon / editor view', None),
    ('204229', 'Aster dragon / full desktop view', '20:42:23'),
    ('204306', 'Petalo / Light and Signalling', '20:42:59'),
    ('204332', 'Explorer Engineer', '20:43:25'),
    ('204404', 'Tiko / Object Manipulation', '20:43:56'),
    ('204425', 'Tolly / Tollgate and Access', '20:44:21'),
    ('204448', 'Quacky / Movement and Delivery', '20:44:43'),
    ('204532', 'Brass lantern', '20:45:26'),
    ('204556', 'Delivery crate and energy crystal', '20:45:52'),
    ('204625', 'Forest tree', '20:46:22'),
    ('204648', 'Fairy cottage', '20:46:41'),
    ('204716', 'Signal crystals', '20:47:12'),
]


def copy_sources():
    DESTINATION.mkdir(exist_ok=True)
    manifest = []
    for stamp, title, clock in SHOTS:
        name = f'Screenshot 2026-09-17 {stamp}.png'
        source = SOURCE / name
        target = DESTINATION / name
        if source.is_file():
            shutil.copy2(source, target)
        if not target.is_file():
            raise FileNotFoundError(source)
        manifest.append({
            'file': name, 'subject': title,
            'visible_date': '17-09-2026' if clock else None,
            'visible_clock': clock,
            'filename_date_time': f'2026-09-17 {stamp[:2]}:{stamp[2:4]}:{stamp[4:]}',
            'sha256': hashlib.sha256(target.read_bytes()).hexdigest(),
        })
    (DESTINATION / 'manifest.json').write_text(json.dumps(manifest, indent=2) + '\n', encoding='utf-8')
    (DESTINATION / 'README.txt').write_text(
        'Twelve participant-supplied Unity Editor screenshots, embedded unchanged in Appendix A.\n'
        'Visible taskbar clocks and filename timestamps are transcribed separately. The first image has no visible taskbar date/time.\n'
        'These images document the asset prefabs in Unity Editor. Gameplay evidence remains in Live-Snaps.\n', encoding='utf-8')
    return manifest


def append_evidence(document):
    if any(para.text == HEADING for para in document.paragraphs):
        raise ValueError('The dated evidence appendix is already present.')
    manifest = copy_sources()
    section = document.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(.35)
    for index, item in enumerate(manifest, 1):
        if index > 1:
            document.add_page_break()
        if index == 1:
            document.add_heading(HEADING, 1)
            note = document.add_paragraph(
                'Participant-supplied screenshots from 17 September 2026. Full images are preserved; captions record the visible clock and the filename timestamp separately.', 'Caption')
            note.paragraph_format.keep_with_next = True
        else:
            document.add_heading(f'Asset evidence A{index:02d} / {item["subject"]}', 1)
        para = document.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.keep_with_next = True
        shape = para.add_run().add_picture(str(DESTINATION / item['file']), width=Inches(9))
        shape._inline.docPr.set('descr', f'Asset evidence A{index:02d}. {item["subject"]}. Original Unity Editor screenshot including all visible desktop details.')
        clock = (f'Visible taskbar: {item["visible_date"]}, {item["visible_clock"]}.'
                 if item['visible_clock'] else 'Taskbar date/time is not visible in this image.')
        caption = document.add_paragraph(
            f'A{index:02d}. {item["subject"]}. {clock}\n'
            f'Original file: {item["file"]} | Filename timestamp: {item["filename_date_time"]}.', 'Caption')
        caption.paragraph_format.keep_together = True
        for run in caption.runs:
            run.font.size = Pt(10)


if __name__ == '__main__':
    source = ROOT / 'Crust-Innovators-Phase-2-Submission.docx'
    output = ROOT / 'Crust-Innovators-Phase-2-Submission-With-Dated-Evidence.docx'
    document = Document(source)
    if any(para.text == HEADING for para in document.paragraphs):
        raise SystemExit('No change: appendix is already present.')
    backup = ROOT.parent / 'QA' / 'DocumentBackups' / 'Submission-before-dated-evidence.docx'
    backup.parent.mkdir(exist_ok=True)
    if not backup.exists():
        shutil.copy2(source, backup)
    append_evidence(document)
    document.save(output)
    print('Added all twelve unchanged screenshots with timestamp captions to the submission.')
