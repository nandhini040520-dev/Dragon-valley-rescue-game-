from pathlib import Path
import re
import shutil
from dated_evidence import append_evidence
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'HackathonDocs'
EVIDENCE = OUT / 'Live-Snaps'
EVIDENCE.mkdir(exist_ok=True)
TEAM = 'Crust Innovators'
NAMES = 'Devanand S, Karthiga S, Dharsana K, Nandhini S'
COLLEGE = 'AVS Engineering College'
INK = '163C36'
MUTED = '5D6E73'
ORANGE = 'C56716'

# compact_reference_guide with named brand-color and evidence-table overrides.
# Letter; 1 in margins; 9360 DXA content width; Calibri 11, line 1.25.
def fmt(style, size, color=INK, before=0, after=6, bold=False, line=1.25):
    style.font.name='Calibri'; style.font.size=Pt(size); style.font.color.rgb=RGBColor.from_string(color); style.font.bold=bold
    p=style.paragraph_format; p.space_before=Pt(before); p.space_after=Pt(after); p.line_spacing=line

def base(title, subtitle):
    d=Document(); s=d.sections[0]
    s.page_width=Inches(8.5); s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)
    s.header_distance=s.footer_distance=Inches(.492)
    fmt(d.styles['Normal'],11)
    fmt(d.styles['Title'],29,INK,0,12,True,1.03)
    fmt(d.styles['Subtitle'],13,MUTED,0,12)
    fmt(d.styles['Heading 1'],16,INK,18,10,True)
    fmt(d.styles['Heading 2'],13,INK,14,7,True)
    fmt(d.styles['Heading 3'],12,ORANGE,10,5,True)
    fmt(d.styles['Caption'],9,MUTED,4,8,False,1.1)
    fmt(d.styles['List Bullet'],11,INK,0,4,False,1.25)
    for style in ['Heading 1','Heading 2','Heading 3','Caption']:
        d.styles[style].paragraph_format.keep_with_next=style!='Caption'
    # Define the bullet indent explicitly rather than inheriting Word defaults.
    numbering=d.part.numbering_part.element
    abstract=OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'),'70')
    level=OxmlElement('w:lvl'); level.set(qn('w:ilvl'),'0')
    for tag,val in [('start','1'),('numFmt','bullet'),('lvlText','•'),('lvlJc','left')]:
        el=OxmlElement('w:'+tag); el.set(qn('w:val'),val); level.append(el)
    props=OxmlElement('w:pPr'); ind=OxmlElement('w:ind'); ind.set(qn('w:left'),'540'); ind.set(qn('w:hanging'),'270'); props.append(ind)
    tabs=OxmlElement('w:tabs'); tab=OxmlElement('w:tab'); tab.set(qn('w:val'),'num'); tab.set(qn('w:pos'),'540'); tabs.append(tab); props.append(tabs); level.append(props); abstract.append(level); numbering.append(abstract)
    num=OxmlElement('w:num'); num.set(qn('w:numId'),'70'); aid=OxmlElement('w:abstractNumId'); aid.set(qn('w:val'),'70'); num.append(aid); numbering.append(num)
    header=s.header.paragraphs[0]; header.text='FUNOBOTZ  |  DISCOVERY WORLD     /     '+TEAM.upper(); fmt(header.style,9,MUTED,0,0,False,1)
    footer=s.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=footer.add_run('PHASE 2  •  '+TEAM+'    |    '); run.font.size=Pt(9); run.font.color.rgb=RGBColor.from_string(MUTED)
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    d.core_properties.title=title; d.core_properties.author=TEAM; d.core_properties.subject=subtitle
    d.add_paragraph(title,'Title'); d.add_paragraph(subtitle,'Subtitle')
    return d

def p(d,text,style=None): return d.add_paragraph(text,style)
def h(d,text): d.add_heading(text,level=1)
def page(d,title): d.add_page_break(); h(d,title)
def bullet(d,text):
    para=d.add_paragraph(text,'List Bullet'); pr=para._p.get_or_add_pPr(); np=OxmlElement('w:numPr')
    for tag,val in [('ilvl','0'),('numId','70')]:
        el=OxmlElement('w:'+tag); el.set(qn('w:val'),val); np.append(el)
    pr.append(np)
def team(d):
    p(d,TEAM+'  |  '+COLLEGE)
    p(d,NAMES)

def table(d,headers,rows,widths,size=9.5):
    t=d.add_table(rows=1,cols=len(headers)); t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.LEFT
    widths=[round(x*1440) for x in widths]; widths[-1]=9360-sum(widths[:-1])
    pr=t._tbl.tblPr
    tw=pr.find(qn('w:tblW')); tw.set(qn('w:w'),'9360'); tw.set(qn('w:type'),'dxa')
    indent=OxmlElement('w:tblInd'); indent.set(qn('w:w'),'120'); indent.set(qn('w:type'),'dxa'); pr.append(indent)
    margins=OxmlElement('w:tblCellMar')
    for edge,n in [('top',80),('bottom',80),('start',120),('end',120)]:
        el=OxmlElement('w:'+edge); el.set(qn('w:w'),str(n)); el.set(qn('w:type'),'dxa'); margins.append(el)
    pr.append(margins)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D7E0DF'); borders.append(el)
    pr.append(borders)
    for col,w in zip(t._tbl.tblGrid.gridCol_lst,widths): col.set(qn('w:w'),str(w))
    for cells,values,is_header in [(t.rows[0].cells,headers,True)]+[(t.add_row().cells,row,False) for row in rows]:
        for idx,(cell,value) in enumerate(zip(cells,values)):
            cell.width=Inches(widths[idx]/1440); cell._tc.get_or_add_tcPr().find(qn('w:tcW')).set(qn('w:w'),str(widths[idx]))
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell.text=str(value)
            for para in cell.paragraphs:
                para.paragraph_format.space_before=Pt(0); para.paragraph_format.space_after=Pt(3); para.paragraph_format.line_spacing=1.08
                for run in para.runs: run.font.size=Pt(size); run.font.bold=is_header; run.font.color.rgb=RGBColor.from_string(INK)
            if is_header:
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E8EEF5'); cell._tc.get_or_add_tcPr().append(shade)
        row=cells[0]._tc.getparent(); rpr=row.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit'); rpr.append(cant)
        if is_header: rep=OxmlElement('w:tblHeader'); rpr.append(rep)
    return t

def image(d,path,caption,width=6.1):
    para=d.add_paragraph(); para.paragraph_format.space_after=Pt(0); para.paragraph_format.keep_with_next=True
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    shape=para.add_run().add_picture(str(path),width=Inches(width)); shape._inline.docPr.set('descr',caption)
    p(d,caption,'Caption')

snap_sources={
 'Snap01-World-Level.png':'AssetPreviews/Regions/2-Hidden-Forest.png',
 'Snap02-Player-Start.png':'QA/Snap02-Player-Start.png',
 'Snap03-Forest-Challenge.png':'QA/Snap03-Forest-Challenge.png',
 'Snap04-Tiko-Lift.png':'QA/Snap04-Tiko-Lift.png',
 'Snap05-Protected-Plants-Feedback.png':'QA/Snap05-Protected-Plants-Feedback.png',
 'Snap06-Forest-Complete.png':'QA/Snap06-Forest-Complete.png',
 'Snap07-Essential-UI.png':'QA/input-05-robot-journal.png',
 'Snap08-Live-Runtime-Structure.png':'QA/Snap08-Live-Runtime-Structure.png',
 'Snap09-Rescue-Prerequisite-Test.png':'QA/Snap09-Rescue-Prerequisite-Test.png',
 'Snap10-Input-Smoke-Evidence.png':'QA/input-04-interact.png',
}

def snap(d,num,caption,width=6.1):
    name=next(n for n in snap_sources if n.startswith(f'Snap{num:02d}-'))
    image(d,EVIDENCE/name,f'Live Snap {num}. '+caption,width)

def submission():
    d=base('Adventure Quest\nThe Dragon Valley Rescue','PHASE 2 PARTICIPANT SUBMISSION  /  FUNOBOTZ')
    team(d)
    image(d,ROOT/'AssetPreviews/Unity-Character-Lineup.png','Actual Unity models: four Funobotz based on the supplied kit references, with the storybook explorer and friendly dragon.',5.7)
    h(d,'Game concept')
    p(d,'Explore seven connected regions as an Explorer Engineer. Observe obstacles, inspect clues, choose the right Funobotz and restore safe routes. The friendly dragon is trapped by a failed restoration system; rescue requires delivery, manipulation, light and access control.')
    p(d,'Theme: Discovery World / Adventure Quest World\nPrimary category: 3D Adventure + Puzzle + Exploration\nDesign target: ages 10–14; child usability still needs external validation.\nPlatform: Windows desktop / Unity 6000.6.0f1 / built-in renderer.')
    p(d,'EXPLORE → OBSERVE → INSPECT → CHOOSE → SELECT FUNOBOTZ → INTERACT → MANIPULATE → COMPLETE → UNLOCK → APPLY', 'Heading 3')
    p(d,'Evidence: gameplay figures show the Unity player; the journal map is declared AI-assisted illustration. This report follows the supplied sample’s structure. Its owl and lever/fulcrum example are outside this prototype.','Caption')

    page(d,'1. 3D level / world map')
    snap(d,1,'Hidden Forest in the live rendered world. The path enters from the near side, the branch challenge occupies the route, and the far checkpoint leads toward Rainbow Bridge.')
    table(d,['Region','Discovery / engineering action','Progression'],[
      ['Grand Gateway','Inspect a route; choose it; Tolly opens access.','Hidden Forest'],
      ['Hidden Forest','Tiko moves the branch to a safe pad; restore marker.','Rainbow Bridge'],
      ['Rainbow Bridge','Align support with Tiko; deliver supply with Quacky.','Castle Place'],
      ['Castle Place','Crystal → container → lift → Petalo receiver.','Mystery Cave'],
      ['Mystery Cave','Petalo reveals markers; Quacky retrieves; light 3 beacons.','Fairy Garden'],
      ['Fairy Garden','Deliver a path kit; place a heavy stone safely.','Dragon Valley'],
      ['Dragon Valley','Restore four systems and free the friendly dragon.','Final portal / victory'],
    ],[1.25,3.8,1.45])
    p(d,'Design notes: the player physically walks the connected trail. Distinct landmarks and visible gates establish progression. Cyan placement pads and pink protected habitats turn spatial observation into a meaningful choice. The complete overhead geometry view is supplied separately in the asset gallery.','Caption')

    page(d,'2. Complete gameplay loop / discover')
    snap(d,2,'Player start and mission entry. A short objective, expedition counter and controls introduce the task.')
    snap(d,3,'Challenge zone. The fallen branch, inspection target and safe placement area are part of the same 3D forest route.')

    page(d,'2. Complete gameplay loop / act and test')
    snap(d,4,'Engineering interaction. Tiko’s segmented arm lifts the inspected branch; the player must choose where to put it.')
    snap(d,5,'Feedback after an unsafe choice. A protected plant bed rejects the placement, shows Environmental Warning and resets the branch for another attempt.')

    page(d,'2. Complete gameplay loop / adjust and unlock')
    snap(d,6,'Completion and progression. The branch has been placed safely, the marker is restored, and the forest mission unlocks Rainbow Bridge.')
    p(d,'Repeatable reasoning loop: observe the blocked route → inspect the branch → select Tiko → lift → compare the marked areas → test a placement → read feedback → lift again if reset → place on the safe foundation → restore the route marker.')
    h(d,'What the player learns')
    bullet(d,'Function matching: a light robot cannot substitute for a manipulator or delivery robot.')
    bullet(d,'Spatial constraints: moving an obstacle is only successful when the destination preserves the protected habitat.')
    bullet(d,'Cause and effect: feedback explains why a placement failed; a safe adjustment changes the world and unlocks a route.')
    p(d,'This prototype approximates engineering through constrained placement, sequencing and system dependencies. It does not simulate lever ratios, torque, structural load or a continuous clearance percentage.','Caption')

    page(d,'3. Essential UI / child usability')
    snap(d,7,'The Q journal shows the illustrated expedition map, the current hint and the seven-region progress list. This is an actual game-UI capture; the map inside it is illustrated artwork.')
    table(d,['Essential state','Purpose and implemented feedback'],[
      ['Start','Start Adventure, Continue, Reset and Quit; clear game title.'],
      ['Mission','One active objective, step count, expedition counter and nearby E / F prompt.'],
      ['Feedback','Short written reason, visual response, sound cue and recoverable object reset.'],
      ['Complete','Mission popup, unlocked route, skill label; final victory offers Play Again.'],
    ],[1.25,5.25])
    bullet(d,'WASD + mouse navigation; E interacts, F inspects, Q opens the journal, 1–4 select robots, Space jumps, Esc pauses.')
    bullet(d,'Text accompanies color cues. Protected beds and safe pads use distinct colors; objects snap into valid targets.')
    p(d,'Usability questions remain: can a fresh player discover F before E, identify the safe pad, and understand how to retry after a reset? Record a real tester’s attempt before claiming child-usability success.','Caption')

    page(d,'4. Technical plan and live structure')
    snap(d,8,'A development-only overlay reads the current runtime objects and state. It is a live player capture, not a mock editor hierarchy.',5.85)
    table(d,['System','Responsibility','Stored state'],[
      ['AdventureGame','World, menus, pause and victory lifecycle.','GameState'],
      ['MissionManager','Objectives, region unlocks, four rescue flags.','Region, step, completed IDs, tokens'],
      ['Explorer / Camera','CharacterController motion and orbit camera.','Velocity / yaw / pitch'],
      ['Interaction framework','Target, inspect, validate and run actions.','Target / inspected / busy / carried'],
      ['RobotTeam + controllers','Select and execute robot-specific actions.','Selected robot / working / held object'],
      ['AdventureWorld','Create zones and apply visible mission changes.','Gates / effects / restoration'],
      ['HUD / AudioManager','Prompts, feedback and generated sound cues.','Notices / UI state / clips'],
    ],[1.5,2.8,2.2],9)
    p(d,'One connected AdventureWorld scene; built-in pipeline, legacy input, Mono Windows x64. PlayerPrefs saves region checkpoints and explorer tokens. Runtime models are authored in code; exported prefabs keep reusable meshes and materials.','Caption')

    page(d,'4. Interaction list')
    table(d,['Object / area','Action and robot','World response'],[
      ['Route signs / Gateway','F inspect, E choose; Tolly at checkpoint.','Route choice recorded; entrance opens.'],
      ['Fallen branch / Forest','Tiko: inspect, lift, place on cyan pad.','Branch leaves route; marker becomes actionable.'],
      ['Protected plants','Attempt placement while carrying.','Environmental Warning; object resets; retry allowed.'],
      ['Bridge support','Tiko lifts and places support in socket.','Missing span aligns into a walkable crossing.'],
      ['Bridge supply','Quacky delivers the supply.','Crossing checkpoint becomes available.'],
      ['Castle checkpoint','Tolly operates gate.','Courtyard access opens.'],
      ['Crystal / container / lift','Collect crystal, load container, Tiko moves it.','Lift transports the loaded container.'],
      ['Tower receiver','Petalo illuminates photosensitive switch.','Hidden passage opens.'],
      ['Cave beacons and passage','Petalo lights markers; Quacky retrieves crystal.','Three beacons complete the cave route.'],
      ['Garden kit / stone','Quacky delivers; Tiko places heavy stone safely.','Route and garden marker restore.'],
      ['Valley stations','Quacky energy; Tiko repair; Petalo beacon; Tolly access.','Four restoration conditions become true.'],
      ['Final restoration altar','E after all four conditions are true.','Barrier disables, dragon emerges, portal and victory activate.'],
    ],[1.65,2.55,2.3],10)
    h(d,'Interaction architecture')
    p(d,'IInteractable → InteractableBase defines reusable inspect/interact behavior. InteractionDetector finds a nearby target; InteractionPrompt and AdventureHUD display it. InteractionManager rejects wrong robots, unmet prerequisites and invalid placement before applying an objective.')

    page(d,'5. Riskiest mechanic prototype')
    p(d,'The key integration risk is completing the rescue too early or allowing state from an earlier attempt to leak into a restart. The final interaction depends on all four robot systems and must survive save/load and reset.')
    snap(d,9,'The acceptance driver removes the energy prerequisite while the other systems are ready. The final rescue is rejected and the game explains the missing conditions.')
    table(d,['Acceptance check','Expected result','Recorded result'],[
      ['Rescue before preparation','Rejected.','PASS'],
      ['Remove energy, repair, beacon or access individually','Each missing system blocks rescue.','PASS'],
      ['All four systems restored','Rescue, emergence, portal and victory.','PASS'],
      ['Unsafe object placement','Reset, re-pick and safe retry work.','PASS'],
      ['Continue from checkpoint','Region and tokens restored; partial puzzle restarts.','PASS'],
      ['Play Again','Objectives and all final flags reset.','PASS'],
    ],[2,3.35,1.15],9.5)
    p(d,'Gate condition: ENERGY_DELIVERED AND MECHANISM_REPAIRED AND BEACON_ACTIVE AND BARRIER_UNLOCKED. Results are implementation checks, not evidence of player comprehension.','Caption')

    page(d,'6. Asset plan and provenance')
    image(d,ROOT/'AssetPreviews/Unity-Character-Lineup.png','Current Unity model lineup. Petalo, Quacky, Tolly and Tiko were corrected against the supplied real-kit views.',6.5)
    table(d,['Asset','Source / license record','Status'],[
      ['Four Funobotz','AI-assisted Unity mesh code using participant-supplied kit references. FUNOBOTZ branding retained; no claim of original brand ownership.','Implemented'],
      ['Reference screenshots','Provided by participant. Used to document physical identity; no independent license document supplied.','Reference only'],
      ['Explorer / dragon / scenery','Custom AI-assisted component models, shared materials and authored environment code.','Implemented'],
      ['UI and puzzle objects','Custom code and geometric shapes. No downloaded marketplace asset pack.','Implemented'],
      ['Illustrated journal map','Built-in image generation; prompt recorded in AssetPreviews/IMAGE-PROMPTS.md.','AI-assisted, declared'],
      ['Sound effects','Generated placeholder tones in AudioManager.','Implemented'],
      ['Segoe UI / Georgia fonts','Local Windows font files; separate redistribution license record not supplied.','Used by prototype'],
    ],[1.35,4,1.15],9)
    p(d,'Reference fidelity: Petalo uses a yellow flower, red/brown smiling LED face and pedestal; Quacky uses a triangular head and flexible red/yellow/green joints; Tolly uses square card panels, dual indicators and a moving barrier; Tiko uses an angular chassis, accordion supports and a segmented gripper. Unseen surfaces and exact dimensions are simplified.','Caption')
    h(d,'24-hour vertical-slice scope')
    p(d,'One connected world, one reusable interaction system, four simple robot actions, seven short missions, a four-system final rescue, menus, checkpoint saves and Windows output. No combat AI, networking, complex inventory, procedural world or advanced physics simulation.')

    page(d,'7. Playtest plan and evidence')
    snap(d,10,'Actual keyboard/mouse smoke test: F inspection followed by E interaction advances the Gateway objective. The tester here is an automated input script; an independent human playtest is still pending.')
    checks=(ROOT/'QA/acceptance-results.txt').read_text(encoding='utf-8-sig')
    count=len(re.findall(r'^PASS ',checks,re.M))
    p(d,f'Current technical evidence: {count} automated acceptance checks passed in the corrected Windows player. Native keyboard/mouse smoke testing covers movement, jump, inspection, interaction, selection, journal, camera, pause, reset and menu return. See the supplied logs for the recorded result.')
    table(d,['Question for a fresh player','Observed evidence','Next action'],[
      ['Understands the mission without explanation?','Not yet tested independently.','Record first action and help requested.'],
      ['Chooses the correct Funobotz?','Functional selection/action checks pass.','Observe unaided function matching.'],
      ['Understands unsafe-placement feedback?','Automated reset/retry check passes.','Record whether the player can recover alone.'],
      ['Completes the loop and follows the unlock?','Full automated route passes.','Run one fresh-player session using the record sheet.'],
    ],[2.2,2.2,2.1],9.5)
    p(d,'Before Phase 3: run an independent session, document the main confusion, make one evidence-based usability change, then retest. Do not substitute the sample’s 18-second observation or other illustrative results.','Caption')

    page(d,'8. Phase 2 self-check against 20 points')
    table(d,['Criterion','Max','Evidence in this submission','Participant check'],[
      ['Core game loop quality','5','Live Snaps 2–6: discover, act, fail, retry, unlock.','Implemented'],
      ['3D experience design','4','Live Snaps 1, 3, 4: navigation and placement.','Implemented'],
      ['Child usability','3','Live Snap 7 + UI decisions; external test pending.','Needs player validation'],
      ['Engineering / problem-solving','3','Snaps 4, 5, 9: constraints and dependencies.','Implemented'],
      ['Technical feasibility','3','Snap 8, Windows build and recorded checks.','Verified technically'],
      ['Visual / interaction clarity','2','Reference-corrected robots, prompts, placement feedback.','Needs fresh-player review'],
    ],[1.65,.4,3,1.45],9)
    p(d,'Maximum available: 20. This is an evidence checklist, not a self-awarded score. Judges determine the score.','Caption')
    table(d,['Live Snap','Evidence file / purpose'],[
      ['1','Snap01-World-Level.png — live 3D forest overview'],
      ['2','Snap02-Player-Start.png — start + objective'],
      ['3','Snap03-Forest-Challenge.png — problem zone'],
      ['4','Snap04-Tiko-Lift.png — core manipulation'],
      ['5','Snap05-Protected-Plants-Feedback.png — failure + reset'],
      ['6','Snap06-Forest-Complete.png — success + progression'],
      ['7','Snap07-Essential-UI.png — actual journal and HUD'],
      ['8','Snap08-Live-Runtime-Structure.png — technical state'],
      ['9','Snap09-Rescue-Prerequisite-Test.png — integration risk'],
      ['10','Snap10-Input-Smoke-Evidence.png — automated input evidence; add human test'],
    ],[.8,5.7],9)
    p(d,'Files are supplied in Live-Snaps. The editable Playtest Record is ready for an external tester. The Windows executable and Unity source remain the technical deliverables. Team identity on this submission is supplied by the participants.','Caption')
    append_evidence(d)
    d.save(OUT/'Crust-Innovators-Phase-2-Submission-With-Dated-Evidence.docx')

def handover():
    d=base('Demo and Technical Handover','ADVENTURE QUEST: THE DRAGON VALLEY RESCUE')
    team(d)
    h(d,'Run the demo')
    p(d,'Extract Adventure-Quest-Windows.zip and launch Adventure Quest.exe. Keep the executable, Data folder, UnityPlayer.dll and MonoBleedingEdge together. Select START ADVENTURE for a clean demonstration; CONTINUE resumes the last saved region checkpoint.')
    table(d,['Control','Action'],[['W A S D','Move'],['Mouse','Orbit the third-person camera'],['Space','Jump'],['F / E','Inspect / interact'],['1 / 2 / 3 / 4','Petalo / Quacky / Tolly / Tiko'],['Q','Illustrated map, hint and region progress'],['Esc','Pause, resume, reset, return to menu or quit']],[1.7,4.8],11)
    h(d,'The first minute')
    p(d,'Introduce the friendly dragon as the rescue goal. Walk toward a route sign, press F to inspect it, then E to choose the route. Select Tolly with 3 and operate the checkpoint. Explain that every robot has a physical function, so the player must choose a suitable tool for each problem.')
    p(d,'Suggested opening: “We are Crust Innovators from AVS Engineering College. This is a peaceful engineering adventure. Observe the obstacle, inspect it, choose the right Funobotz, and restore the path. The final valley needs all four robot systems working together.”')

    page(d,'Seven-region demonstration route')
    for title,text in [
      ('1. Grand Gateway','Inspect either route sign, choose it, then use Tolly at the gateway checkpoint. Both routes join the same expedition.'),
      ('2. Hidden Forest','Inspect the branch. Select Tiko, lift it, and demonstrate a protected pink bed once. Read the warning, lift the reset branch, place it on the cyan pad, and restore the route marker.'),
      ('3. Rainbow Bridge','Inspect the survey post, lift the support with Tiko and place it in the near-bank socket. Use Quacky for the supply delivery. Walk across to the crossing checkpoint.'),
      ('4. Castle Place','Tolly opens the courtyard. Collect the crystal, put it in the container, use Tiko to move the container onto the lift, then Petalo activates the tower light receiver.'),
      ('5. Mystery Cave','Inspect and illuminate the marker stone with Petalo. Light the first beacon, use Quacky for the narrow passage retrieval, then activate the remaining two beacons.'),
      ('6. Fairy Garden','Quacky delivers the path kit. Inspect and lift the heavy stone with Tiko. Place it on the cyan foundation and activate the garden marker.'),
      ('7. Dragon Valley','Use Quacky for energy, Tiko for mechanical alignment, Petalo for the beacon and Tolly for access. Read the four system indicators, then press E at the restoration altar. Show dragon emergence, victory and Play Again.'),
    ]: d.add_heading(title,2); p(d,text)
    p(d,'This is a suggested walkthrough, not a measured completion-time claim. Use the Q journal whenever the current hint is needed.','Caption')

    page(d,'Technical handover and release checks')
    table(d,['Location','Purpose'],[
      ['Assets/Scenes/AdventureWorld.unity','Single connected-world scene. Runtime director builds the authored level on Play.'],
      ['Assets/Scripts/MissionManager.cs','Mission definitions, objective order and rescue prerequisites.'],
      ['Assets/Scripts/FunobotzModels.cs','Reference-based kit geometry and printed identity.'],
      ['Assets/Scripts/FunobotzMotion.cs','Reusable tollgate, segmented-arm and crawling-joint motion.'],
      ['Assets/Art/Prefabs','Exported reusable models and referenced assets.'],
      ['Build-Windows.ps1 / Run-QA.ps1','Build helper / opt-in standalone acceptance run.'],
      ['QA/acceptance-results.txt','Actual pass/fail results; never infer a pass from the presence of tests.'],
      ['AssetPreviews/index.html','Actual model and region images, plus separately labelled illustration.'],
    ],[2.8,3.7],9.5)
    h(d,'Implementation choices')
    bullet(d,'Unity 6000.6.0f1, built-in renderer, Mono Windows x64. No external online service is required.')
    bullet(d,'PlayerPrefs stores region progress and tokens. An unfinished region restarts from its checkpoint when continued.')
    bullet(d,'Robot movement and object placement are intentionally authored and approximate. Physics realism is outside this prototype’s scope.')
    bullet(d,'Real-kit screenshots guide the Funobotz silhouettes; the storybook environment and explorer remain original prototype art.')
    h(d,'Before a submission upload')
    bullet(d,'Run the extracted game, verify a clean start, and keep the full Windows folder together.')
    bullet(d,'Review the Phase 2 document and add the external player observation to the Playtest Record.')
    bullet(d,'Use actual live snaps as evidence. Declare the illustrated map and AI-assisted asset/code workflow.')
    d.save(OUT/'Crust-Innovators-Demo-and-Handover.docx')

def playtest():
    d=base('External Playtest Record','BLANK RECORDING FORM  /  To be completed after an actual player session')
    team(d)
    p(d,'Game: Adventure Quest: The Dragon Valley Rescue\nSession date/time: ____________________    Observer: ____________________\nTester alias: ____________________    Age range: ____________________\nBuild / device: ____________________    Previous experience: ____________________')
    h(d,'Run the session')
    p(d,'Ask the player to restore the Hidden Forest path without explaining the solution. Let them read the controls and journal. Record confusion and help requested before giving a hint. Use an alias; do not include unnecessary personal details in the submission.')
    table(d,['Task','Outcome / observations','Help needed?'],[
      ['Start and understand the objective','First action: __________________\nObservation: __________________','Yes / No\nHint: ______'],
      ['Inspect the fallen branch','Found F independently? _________\nObservation: __________________','Yes / No\nHint: ______'],
      ['Choose the correct Funobotz','Choice and reason: _____________\nObservation: __________________','Yes / No\nHint: ______'],
      ['Place safely or recover from a warning','Attempt / adjustment: __________\nObservation: __________________','Yes / No\nHint: ______'],
      ['Restore marker and follow the unlock','Recognized completion? __________\nObservation: __________________','Yes / No\nHint: ______'],
    ],[1.7,3.6,1.2],10)
    p(d,'Completion time (measure it): ______________    Completed unaided? Yes / No\nEvidence screenshot / recording filename: ______________________________\nPlayer’s own explanation of the solution: _______________________________\n________________________________________________________________')
    p(d,'No results have been pre-filled. Automated acceptance tests do not count as an external player session.','Caption')
    page(d,'Observation → design change → retest')
    h(d,'Main usability finding')
    p(d,'What was the most important confusion?\n________________________________________________________________\n________________________________________________________________\nWhat did the player do or say that supports this finding?\n________________________________________________________________\n________________________________________________________________')
    h(d,'One change before Phase 3')
    p(d,'Change to the UI, clue, interaction or layout:\n________________________________________________________________\n________________________________________________________________\nWhy should it help?\n________________________________________________________________\nOwner: ____________________    Implemented date: ____________________')
    table(d,['Measure','First session','Retest'],[
      ['Understands the objective without help','________________','________________'],
      ['Finds inspect / interaction controls','________________','________________'],
      ['Chooses the correct robot','________________','________________'],
      ['Recovers from failure feedback','________________','________________'],
      ['Completes the loop','________________','________________'],
    ],[3,1.75,1.75],10)
    p(d,'Retest date / tester alias: __________________________________________\nDecision: Keep change / Revise again\nRemaining issue: __________________________________________________')
    d.save(OUT/'Crust-Innovators-Playtest-Record.docx')

if __name__=='__main__':
    for name,source in snap_sources.items():
        assert (ROOT/source).is_file(),f'Missing current evidence: {source}'
        shutil.copy2(ROOT/source,EVIDENCE/name)
    assert 'RESULT: PASSED' in (ROOT/'QA/acceptance-results.txt').read_text(encoding='utf-8-sig')
    submission(); handover(); playtest()
    (OUT/'Live-Snaps/README.txt').write_text('Actual Unity captures from the corrected prototype. Snap 8 is a live runtime development overlay. Snap 10 is automated native input evidence; an independent player observation remains pending. No illustrative sample screenshots are used.\n',encoding='utf-8')
    print('Created three editable hackathon documents and copied the ten live evidence images.')
